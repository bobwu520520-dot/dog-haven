# -*- coding: utf-8 -*-
"""
《金毛和汤》数值配置与经济平衡表生成器
--------------------------------------------------
数值单一数据源：本脚本不手抄任何数值，全部读取 scratch/_config_dump.json
（该文件由 scratch/dump_config.js 从 js/config.js 直接导出）。
用法：
    node scratch/dump_config.js
    python generate_excel.py
"""
import json
import math
import os

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DUMP_PATH = os.path.join(BASE_DIR, "scratch", "_config_dump.json")
OUT_PATH = os.path.join(BASE_DIR, "金毛和汤_数值配置与经济平衡表.xlsx")

if not os.path.exists(DUMP_PATH):
    raise SystemExit(
        "缺少配置快照 scratch/_config_dump.json\n"
        "请先执行：  node scratch/dump_config.js\n"
        "该快照由 js/config.js 直接导出，保证平衡表与游戏代码数值永远一致。"
    )

with open(DUMP_PATH, "r", encoding="utf-8") as f:
    CFG = json.load(f)

FACILITIES = CFG["FACILITIES_CONFIG"]
CUISINES = CFG["DOG_CUISINES_CONFIG"]
DOGS = CFG["DOG_BREEDS_CONFIG"]
OUTFITS = CFG["OUTFITS_CONFIG"]
COATS = CFG["COAT_COLORS_CONFIG"]
STAGES = CFG["TOWN_EXPANSION_STAGES"]
RECIPE = CFG["RECIPE_UPGRADE_CONFIG"]
ACHIEVEMENTS = CFG["ACHIEVEMENTS_CONFIG"]
DAILY_TASKS = CFG["DAILY_TASKS_CONFIG"]
WEATHER = CFG["WEATHER_CONFIG"]
TOYS = CFG["PLAY_TOYS_CONFIG"]
PARK = CFG["PARK_ACTIVITIES_CONFIG"]

MAX_LEVEL = 50
AUTO_COLLECT_LEVEL = 20

# ---------------- 样式 ----------------
F_TITLE = Font(name="微软雅黑", size=14, bold=True, color="2C3E50")
F_HEAD = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
F_BOLD = Font(name="微软雅黑", size=10, bold=True, color="2C3E50")
F_TEXT = Font(name="微软雅黑", size=10, color="2C3E50")
F_ACCENT = Font(name="微软雅黑", size=10, bold=True, color="D35400")

FILL_HEAD = PatternFill("solid", start_color="E67E22")
FILL_SUB = PatternFill("solid", start_color="FAD7A0")
FILL_ACCENT = PatternFill("solid", start_color="FEF9E7")
FILL_OK = PatternFill("solid", start_color="EAFAF1")

BORDER = Border(*[Side(style="thin", color="E5E7E9")] * 4)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def add_sheet(wb, title):
    ws = wb.create_sheet(title=title)
    ws.views.sheetView[0].showGridLines = True
    return ws


def write_title(ws, text, ncols, height=34):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws["A1"].alignment = CENTER
    ws.row_dimensions[1].height = height


def write_header(ws, row, headers):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill = FILL_HEAD
        c.font = F_HEAD
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[row].height = 26


def write_row(ws, row, values, fills=None, bold=False):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = F_BOLD if bold else F_TEXT
        c.border = BORDER
        c.alignment = CENTER if (not isinstance(v, str) or len(str(v)) < 22) else LEFT
        if fills and fills.get(i):
            c.fill = fills[i]


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def upgrade_cost(base_cost, multiplier, level):
    """与 WangwangFormulas.getFacilityUpgradeCost 一致：base * mult^level"""
    return math.floor(base_cost * (multiplier ** level))


def price_at_level(station, level):
    """与 kitchen.calcActualDishPrice 的基础部分一致：dish.basePrice * (1 + (lv-1)*0.10)"""
    idx = 2 if level >= 25 else (1 if level >= 10 else 0)
    dish = station["dishes"][idx]
    return idx, dish, round(dish["basePrice"] * (1 + (level - 1) * 0.10))


# ============================================================
# Sheet 1 · 5 大料理工位升级与菜品产出曲线
# ============================================================
def sheet_facility_curve(wb):
    ws = add_sheet(wb, "料理工位升级曲线")
    headers = ["料理工位", "所属菜系", "等级", "升级费用(🪙)", "累计投入(🪙)", "制作耗时(秒)",
               "当前料理", "基础售价(🪙)", "实际售价(🪙)", "单工位产出(🪙/秒)", "自动收取"]
    write_title(ws, "《金毛和汤》5 大料理工位升级与三阶料理产出曲线 (Lv.1 ~ Lv.50)", len(headers))
    write_header(ws, 2, headers)

    row = 3
    for fid, st in FACILITIES.items():
        cumulative = st["unlockCost"]
        for lv in range(1, MAX_LEVEL + 1):
            if lv > 1:
                cumulative += upgrade_cost(st["baseUpgradeCost"], st["costMultiplier"], lv - 1)
            cost = 0 if lv == 1 else upgrade_cost(st["baseUpgradeCost"], st["costMultiplier"], lv - 1)
            idx, dish, price = price_at_level(st, lv)
            speed_reduction = min(0.8, (lv - 1) * st["speedPerLevel"])
            cook_time = round(dish["baseTime"] * (1 - speed_reduction), 2)
            gps = round(price / cook_time, 2) if cook_time else 0
            auto = "✅ 已解锁" if lv >= AUTO_COLLECT_LEVEL else "—"
            write_row(ws, row, [
                st["name"], CUISINES[st["cuisineId"]]["name"], lv, cost, cumulative,
                cook_time, dish["name"], dish["basePrice"], price, gps, auto
            ], fills={11: FILL_OK if lv >= AUTO_COLLECT_LEVEL else None})
            row += 1
        row += 1
    autosize(ws, [14, 12, 8, 14, 14, 12, 20, 12, 12, 15, 10])
    ws.freeze_panes = "C3"
    return ws


# ============================================================
# Sheet 2 · 5 大料理菜系与食谱升级
# ============================================================
def sheet_cuisine(wb):
    ws = add_sheet(wb, "料理菜系与食谱")
    headers = ["菜系", "图标", "情感标签", "解锁小镇等级", "产出工位", "食谱满级售价加成", "三阶代表料理"]
    write_title(ws, "《金毛和汤》5 大狗狗料理菜系与食谱升级 (GDD 1.2)", len(headers))
    write_header(ws, 2, headers)

    row = 3
    for cid, c in CUISINES.items():
        fac = FACILITIES.get(c["facilityId"], {})
        max_bonus = f"+{round((RECIPE['maxLevel'] - 1) * RECIPE['priceBonusPerLevel'] * 100)}%"
        dishes = " / ".join(d["name"] for d in c["dishes"])
        write_row(ws, row, [c["name"], c["icon"], c["tag"], f"Lv.{c['stage']}",
                            f"{fac.get('icon','')} {fac.get('name','')}", max_bonus, dishes])
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
    c = ws.cell(row=row, column=1, value="食谱升级曲线 (升级食谱 → 对应料理工位售价提升)")
    c.font = F_ACCENT
    c.fill = FILL_SUB
    row += 1
    write_header(ws, row, ["食谱等级", "升级费用(🪙)", "累计投入(🪙)", "该菜系售价加成"])
    row += 1
    cumulative = 0
    for lv in range(1, RECIPE["maxLevel"] + 1):
        cost = 0 if lv == 1 else math.floor(RECIPE["baseCost"] * (RECIPE["costMultiplier"] ** (lv - 2)))
        if lv > 1:
            cumulative += cost
        bonus = f"+{round((lv - 1) * RECIPE['priceBonusPerLevel'] * 100)}%"
        write_row(ws, row, [f"Lv.{lv}", cost, cumulative, bonus])
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
    c = ws.cell(row=row, column=1,
                value=f"5 大食谱全部满级额外获得 +{round(RECIPE['maxLevelGlobalSpeedBonus']*100)}% 全镇产速加成")
    c.font = F_ACCENT
    c.fill = FILL_ACCENT
    autosize(ws, [12, 8, 12, 16, 20, 20, 46])
    return ws


# ============================================================
# Sheet 3 · 狗狗主厨名册与技能
# ============================================================
def sheet_dogs(wb):
    ws = add_sheet(wb, "狗狗主厨名册")
    headers = ["犬种", "头衔", "职业", "招募费用(🪙)", "专属技能", "技能类型",
               "绑定工位", "基础加成", "满好感加成(Lv.10)", "性格描述"]
    write_title(ws, "《金毛和汤》8 大狗狗主厨名册与专属技能 (GDD 1.2 犬种图鉴)", len(headers))
    write_header(ws, 2, headers)

    TYPE_CN = {
        "local_price": "本工位售价", "local_speed": "本工位产速", "global_speed": "全局产速",
        "collect_speed": "全局自动收取", "offline_limit": "离线收益上限",
        "high_tier_rate": "高级料理概率", "affection_boost": "好感获取"
    }
    row = 3
    for did, d in DOGS.items():
        fac = FACILITIES.get(d.get("targetFacility") or "", {})
        fac_txt = f"{fac.get('icon','')} {fac.get('name','')}" if fac else "— 全局辅助"
        base = d["baseSkillBonus"]
        if d["skillType"] == "offline_limit":
            base_txt, max_txt = f"+{base} 小时", f"+{round(base*1.9,1)} 小时"
        elif d["skillType"] in ("local_speed", "collect_speed"):
            base_txt, max_txt = f"-{round(base*100)}%", f"-{round(base*1.9*100)}%"
        else:
            base_txt, max_txt = f"+{round(base*100)}%", f"+{round(base*1.9*100)}%"
        write_row(ws, row, [d["name"], d["title"], d["profession"], d["recruitCost"],
                            d["skillName"], TYPE_CN.get(d["skillType"], d["skillType"]),
                            fac_txt, base_txt, max_txt, d["personality"]])
        row += 1
    autosize(ws, [10, 16, 12, 14, 14, 16, 16, 14, 18, 52])
    return ws


# ============================================================
# Sheet 4 · 小镇 5 大阶段开荒
# ============================================================
def sheet_stages(wb):
    ws = add_sheet(wb, "小镇开荒阶段")
    headers = ["阶段", "小镇形态", "解锁料理工位", "解锁狗狗", "解锁乐园玩具",
               "开荒条件", "奖励骨头(🦴)", "阶段描述"]
    write_title(ws, "《金毛和汤》小镇 5 大演化阶段与开荒条件 (GDD 1.3 外圈长线)", len(headers))
    write_header(ws, 2, headers)

    row = 3
    for s in STAGES:
        facs = " / ".join(f"{FACILITIES[f]['icon']}{FACILITIES[f]['name']}"
                          for f in s["unlockedFacilities"] if f in FACILITIES)
        dogs = " / ".join(DOGS[d]["name"] for d in s["unlockedDogs"] if d in DOGS)
        toys = " / ".join(t["icon"] + t["name"] for t in TOYS if t["id"] in s["unlockedToys"]) or "—"
        write_row(ws, row, [f"Stage {s['stage']}", s["name"], facs, dogs, toys,
                            s["requirements"]["desc"], s["rewardBones"], s["desc"]])
        row += 1
    autosize(ws, [10, 18, 40, 34, 26, 34, 14, 54])
    return ws


# ============================================================
# Sheet 5 · 收集系统 (服装 + 毛色)
# ============================================================
def sheet_collection(wb):
    ws = add_sheet(wb, "收集系统")
    headers = ["类别", "部位/犬种", "名称", "图标", "货币", "价格", "说明"]
    write_title(ws, "《金毛和汤》收集三要素：犬种 × 毛色 × 饰品 (GDD 1.2 收集亮点)", len(headers))
    write_header(ws, 2, headers)

    row = 3
    for o in OUTFITS:
        part = {"hat": "帽子", "cloth": "衣服", "acc": "配饰"}.get(o["type"], o["type"])
        cur = "🪙 金币" if o["costType"] == "gold" else "🦴 骨头"
        write_row(ws, row, ["服装", part, o["name"], o["icon"], cur, o["cost"], o["desc"]])
        row += 1

    for did, lst in COATS.items():
        breed = DOGS.get(did, {}).get("name", did)
        for c in lst:
            cur = {"gold": "🪙 金币", "bone": "🦴 骨头", "free": "免费"}.get(c["costType"], c["costType"])
            write_row(ws, row, ["毛色", breed, c["name"], "🎨", cur, c["cost"], c["desc"]])
            row += 1

    row += 1
    total_coats = sum(len(v) for v in COATS.values())
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
    c = ws.cell(row=row, column=1,
                value=f"合计：{len(OUTFITS)} 套服装 + {total_coats} 款毛色 = {len(OUTFITS)+total_coats} 项可收集内容")
    c.font = F_ACCENT
    c.fill = FILL_ACCENT
    autosize(ws, [8, 12, 20, 8, 12, 10, 46])
    ws.freeze_panes = "A3"
    return ws


# ============================================================
# Sheet 6 · 任务、成就、天气与迷你玩法
# ============================================================
def sheet_tasks(wb):
    ws = add_sheet(wb, "任务成就与天气")
    headers = ["类型", "名称", "目标/数值", "奖励骨头(🦴)", "说明"]
    write_title(ws, "《金毛和汤》每日任务 / 永久成就 / 天气与迷你玩法", len(headers))
    write_header(ws, 2, headers)

    row = 3
    for t in DAILY_TASKS:
        write_row(ws, row, ["每日任务", t["name"], t["target"], t["rewardBones"], t["desc"]])
        row += 1
    for a in ACHIEVEMENTS:
        write_row(ws, row, ["永久成就", a["name"], a["target"], a["rewardBones"], a["desc"]])
        row += 1

    row += 1
    write_header(ws, row, ["类型", "名称", "数值", "", "说明"])
    row += 1
    for wid, w in WEATHER.items():
        write_row(ws, row, ["天气", f"{w['icon']} {w['name']}", f"产速 +{round(w['gpsBonus']*100)}%", "", w["desc"]])
        row += 1

    row += 1
    write_header(ws, row, ["类型", "名称", "冷却(秒)", "好感奖励", "奖励说明"])
    row += 1
    for pid, p in PARK.items():
        cd = p.get("cooldown", "—")
        aff = p.get("rewardAffection", "—")
        desc = p.get("desc", "")
        if "rewards" in p:
            desc = " / ".join(f"{r['desc']}(权重{r['weight']})" for r in p["rewards"])
        write_row(ws, row, ["迷你玩法", f"{p.get('icon','')} {p['name']}", cd, aff, desc])
        row += 1

    row += 1
    write_header(ws, row, ["类型", "名称", "效果", "", "说明"])
    row += 1
    for t in TOYS:
        write_row(ws, row, ["乐园玩具", f"{t['icon']} {t['name']}", f"坐标({t['x']},{t['y']})", "", t["desc"]])
        row += 1
    autosize(ws, [12, 16, 16, 16, 60])
    return ws


def gps_at(station, level):
    """工位在指定等级的单台产出（🪙/秒），与 kitchen.calcActualCookTime/calcActualDishPrice 基础口径一致"""
    _, dish, price = price_at_level(station, level)
    sr = min(0.8, (level - 1) * station["speedPerLevel"])
    t = dish["baseTime"] * (1 - sr)
    return price / t if t else 0.0


# ============================================================
# Sheet 7 · 平衡性分析
# ============================================================
def sheet_balance(wb):
    ws = add_sheet(wb, "平衡性分析")
    headers = ["料理工位", "解锁费用(🪙)", "Lv.1 产出(🪙/s)", "Lv.10 产出", "Lv.25 产出", "Lv.50 产出",
               "成长倍数", "边际回本 Lv.5", "边际回本 Lv.10", "边际回本 Lv.25", "边际回本 Lv.50"]
    write_title(ws, "《金毛和汤》5 大料理工位经济平衡性分析 (不含狗狗/食谱/天气加成)", len(headers))
    write_header(ws, 2, headers)

    row = 3
    for fid, st in FACILITIES.items():
        marginal = {}
        for L in (5, 10, 25, 50):
            c = upgrade_cost(st["baseUpgradeCost"], st["costMultiplier"], L - 1)
            gain = gps_at(st, L) - gps_at(st, L - 1)
            marginal[L] = (c / gain) if gain > 0 else float("inf")

        def fmt(sec):
            if sec == float("inf"):
                return "—"
            if sec < 90:
                return f"{sec:.0f} 秒"
            if sec < 5400:
                return f"{sec/60:.1f} 分"
            return f"{sec/3600:.1f} 小时"

        g1, g10, g25, g50 = (gps_at(st, L) for L in (1, 10, 25, 50))
        write_row(ws, row, [
            f"{st['icon']} {st['name']}", st["unlockCost"],
            round(g1, 2), round(g10, 2), round(g25, 2), round(g50, 2),
            f"{g50/g1:.1f}×",
            fmt(marginal[5]), fmt(marginal[10]), fmt(marginal[25]), fmt(marginal[50])
        ])
        row += 1

    row += 2
    notes = [
        "口径说明：",
        "· 产出 = 实际售价 ÷ 制作耗时，仅含「工位等级 +10%/级售价、−4%/级耗时(上限 −80%)」，不含狗狗技能、食谱加成、天气加成与 2 倍加速。",
        "· 边际回本时长 = 升到该级的花费 ÷ 该级带来的每秒产出增量。这是判断「升级是否值得点」的核心指标，数值应随等级单调递增。",
        "· 健康曲线：边际回本时长应平滑递增（例如 Lv.5 约 1 分钟 → Lv.25 约 10 分钟 → Lv.50 约 1 小时），既让玩家早期爽快点升级，又让后期需要攒钱。",
        "· 已发现的问题：① Lv.10 / Lv.25 因解锁高阶料理出现「回本断崖」，边际回本反而低于前一级；② 各工位解锁费用相对当前产出能力偏低，新工位几乎瞬间回本。",
    ]
    for n in notes:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        c = ws.cell(row=row, column=1, value=n)
        c.font = F_BOLD if n == notes[0] else F_TEXT
        c.alignment = LEFT
        row += 1
    autosize(ws, [16, 14, 14, 12, 12, 12, 12, 15, 15, 15, 15])
    return ws


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    sheet_facility_curve(wb)
    sheet_cuisine(wb)
    sheet_dogs(wb)
    sheet_stages(wb)
    sheet_collection(wb)
    sheet_tasks(wb)
    sheet_balance(wb)

    wb.save(OUT_PATH)
    print(f"已生成: {OUT_PATH}")
    print(f"  工作表: {', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()
